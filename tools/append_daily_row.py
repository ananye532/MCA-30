#!/usr/bin/env python3
"""Append (or update) one day's row in the Daily Log sheet of the tracker workbook.

Usage:
    python3 tools/append_daily_row.py --entry pending/2026-09-21.json
    python3 tools/append_daily_row.py --date 2026-09-21 --sleep 480 --fitness 30 \
        --study 120 --coding 45 --class-min 250 --classes 5 --other 60 \
        --feeling Good --satisfaction Satisfied --energy High --notes "..."

The row is written in the workbook's own layout: columns I and J stay formulas,
styles are copied from the previous data row, and the existing dropdown
validations (K/L/M) continue to apply. Re-running for a date that is already
present overwrites that row instead of adding a duplicate.
"""

import argparse
import datetime as dt
import json
import sys
from copy import copy
from pathlib import Path

import openpyxl

SHEET = "Daily Log"
FIRST_DATA_ROW = 6
LAST_DATA_ROW = 401

FEELINGS = ["Excellent", "Good", "Neutral", "Low", "Stressed"]
SATISFACTIONS = [
    "Very Satisfied",
    "Satisfied",
    "Neutral",
    "Unsatisfied",
    "Very Unsatisfied",
]
ENERGIES = ["High", "Medium", "Low"]

# Ceilings a row in the sheet may not exceed.
MAX_CLASSES = 8
MAX_SLEEP = 599  # keep sleep under 600 minutes
MAX_FITNESS = 40  # physical activity never above 40 minutes

# Column order of the Daily Log sheet.
COL = {
    "date": 1,
    "sleep": 2,
    "fitness": 3,
    "study": 4,
    "coding": 5,
    "class_min": 6,
    "classes": 7,
    "other": 8,
    "total": 9,
    "free": 10,
    "feeling": 11,
    "satisfaction": 12,
    "energy": 13,
    "notes": 14,
}

REQUIRED = [
    "sleep",
    "fitness",
    "study",
    "coding",
    "class_min",
    "classes",
    "other",
    "feeling",
    "satisfaction",
    "energy",
]


def parse_date(value):
    if isinstance(value, dt.datetime):
        return value
    if isinstance(value, dt.date):
        return dt.datetime(value.year, value.month, value.day)
    return dt.datetime.strptime(str(value).strip(), "%Y-%m-%d")


def find_rows(ws):
    """Return (last_used_row, {date -> row})."""
    last = FIRST_DATA_ROW - 1
    by_date = {}
    for row in range(FIRST_DATA_ROW, LAST_DATA_ROW + 1):
        value = ws.cell(row, COL["date"]).value
        if value is None:
            continue
        last = row
        try:
            by_date[parse_date(value).date()] = row
        except (ValueError, TypeError):
            pass
    return last, by_date


def copy_row_style(ws, src_row, dst_row):
    for col in range(1, len(COL) + 1):
        src = ws.cell(src_row, col)
        dst = ws.cell(dst_row, col)
        dst._style = copy(src._style)
    ws.row_dimensions[dst_row].height = ws.row_dimensions[src_row].height


def validate(entry):
    missing = [k for k in REQUIRED if entry.get(k) in (None, "")]
    if missing:
        sys.exit(f"error: missing values for {', '.join(missing)}")
    for key, allowed in (
        ("feeling", FEELINGS),
        ("satisfaction", SATISFACTIONS),
        ("energy", ENERGIES),
    ):
        if entry[key] not in allowed:
            sys.exit(f"error: {key} must be one of {allowed}, got {entry[key]!r}")
    for key in ("sleep", "fitness", "study", "coding", "class_min", "classes", "other"):
        value = entry[key]
        if not isinstance(value, int) or value < 0:
            sys.exit(f"error: {key} must be a non-negative whole number, got {value!r}")
    if entry["classes"] > MAX_CLASSES:
        sys.exit(
            f"error: classes must be at most {MAX_CLASSES}, got {entry['classes']}"
        )
    if entry["sleep"] > MAX_SLEEP:
        sys.exit(
            f"error: sleep must be at most {MAX_SLEEP} minutes, got {entry['sleep']}"
        )
    if entry["fitness"] > MAX_FITNESS:
        sys.exit(
            f"error: fitness must be at most {MAX_FITNESS} minutes, "
            f"got {entry['fitness']}"
        )
    total = sum(
        entry[k] for k in ("sleep", "fitness", "study", "coding", "class_min", "other")
    )
    if total > 1440:
        sys.exit(f"error: tracked minutes total {total}, more than the 1440 in a day")


def write_row(ws, row, entry, style_src):
    if style_src is not None and style_src != row:
        copy_row_style(ws, style_src, row)
    ws.cell(row, COL["date"]).value = entry["date"]
    for key in ("sleep", "fitness", "study", "coding", "class_min", "classes", "other"):
        ws.cell(row, COL[key]).value = entry[key]
    ws.cell(row, COL["total"]).value = (
        f'=IF(SUM(B{row}:F{row},H{row})=0,"",SUM(B{row}:F{row},H{row}))'
    )
    ws.cell(row, COL["free"]).value = f'=IF(I{row}="","",1440-I{row})'
    ws.cell(row, COL["feeling"]).value = entry["feeling"]
    ws.cell(row, COL["satisfaction"]).value = entry["satisfaction"]
    ws.cell(row, COL["energy"]).value = entry["energy"]
    ws.cell(row, COL["notes"]).value = entry.get("notes", "")


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument(
        "--workbook",
        default=str(Path(__file__).resolve().parent.parent / "mca31.xlsx"),
        help="workbook to edit (default: mca31.xlsx at the repo root)",
    )
    ap.add_argument("--entry", help="JSON file holding the day's values")
    ap.add_argument("--date", help="YYYY-MM-DD (default: today)")
    ap.add_argument("--sleep", type=int)
    ap.add_argument("--fitness", type=int)
    ap.add_argument("--study", type=int)
    ap.add_argument("--coding", type=int)
    ap.add_argument("--class-min", type=int, dest="class_min")
    ap.add_argument("--classes", type=int)
    ap.add_argument("--other", type=int)
    ap.add_argument("--feeling", choices=FEELINGS)
    ap.add_argument("--satisfaction", choices=SATISFACTIONS)
    ap.add_argument("--energy", choices=ENERGIES)
    ap.add_argument("--notes", default="")
    ap.add_argument(
        "--if-exists",
        choices=["update", "skip", "error"],
        default="update",
        help="what to do when the date already has a row (default: update it)",
    )
    args = ap.parse_args()

    entry = {}
    if args.entry:
        entry.update(json.loads(Path(args.entry).read_text()))
    for key in REQUIRED + ["date", "notes"]:
        value = getattr(args, key, None)
        if value not in (None, ""):
            entry[key] = value
    entry["date"] = parse_date(entry.get("date") or dt.date.today())
    validate(entry)

    path = Path(args.workbook)
    wb = openpyxl.load_workbook(path)
    if SHEET not in wb.sheetnames:
        sys.exit(f"error: {path} has no {SHEET!r} sheet")
    ws = wb[SHEET]

    last, by_date = find_rows(ws)
    target = by_date.get(entry["date"].date())
    if target is None:
        target = last + 1
        if target > LAST_DATA_ROW:
            sys.exit(f"error: sheet is full at row {LAST_DATA_ROW}")
        action = "added"
    elif args.if_exists == "skip":
        print(
            f"skipped: {path.name} already has row {target} for "
            f"{entry['date'].date().isoformat()}"
        )
        return
    elif args.if_exists == "error":
        sys.exit(
            f"error: {path.name} already has row {target} for "
            f"{entry['date'].date().isoformat()}"
        )
    else:
        action = "updated"
    write_row(ws, target, entry, style_src=last if last >= FIRST_DATA_ROW else None)
    wb.save(path)
    print(
        f"{action} row {target} of {path.name} for "
        f"{entry['date'].date().isoformat()}"
    )


if __name__ == "__main__":
    main()
